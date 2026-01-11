"""Модуль для работы с Excel-файлами"""
import logging
from pathlib import Path
from typing import List, Dict, Any
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill
import pytz
from config.settings import EXCEL_FILE_PATH, TIMEZONE, SHEET_NAMES

logger = logging.getLogger(__name__)


class ExcelManager:
    """Менеджер для работы с Excel-файлами"""

    def __init__(self):
        """Инициализация менеджера Excel"""
        self.excel_file = Path(EXCEL_FILE_PATH)
        self.sheets = SHEET_NAMES

        # Создаём файл если не существует
        if not self.excel_file.exists():
            self._create_excel_file()

        logger.info(f"Excel менеджер инициализирован: {self.excel_file}")

    def _create_excel_file(self):
        """Создать Excel файл с необходимыми листами"""
        try:
            wb = openpyxl.Workbook()
            wb.remove(wb.active)  # Удаляем дефолтный лист

            # Создаём листы
            for sheet_key, sheet_name in self.sheets.items():
                ws = wb.create_sheet(sheet_name)

                # Добавляем заголовки
                headers = self._get_headers(sheet_key)
                ws.append(headers)

                # Форматируем заголовки
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

            wb.save(self.excel_file)
            logger.info(f"Создан новый Excel файл: {self.excel_file}")

        except Exception as e:
            logger.error(f"Ошибка при создании Excel файла: {e}")
            raise

    def _get_headers(self, sheet_key: str) -> List[str]:
        """Получить заголовки для листа"""
        headers_map = {
            'queries': ['ID', 'Дата создания', 'Тема', 'Город', 'Запрос', 'Статус', 'Дата обработки'],
            'sites': ['ID', 'Дата добавления', 'Название компании', 'Сайт', 'Телефон',
                     'Адрес', 'Город', 'Категория', 'Статус сбора данных', 'ID запроса'],
            'companies': ['ID', 'Дата сбора', 'Название компании', 'Сайт', 'Email',
                         'Телефон', 'Адрес', 'Описание', 'Услуги', 'Контактное лицо',
                         'Соц. сети', 'ID сайта'],
            'letters': ['ID', 'Дата создания', 'Название компании', 'Email',
                       'Тема письма', 'Текст письма', 'Статус отправки', 'ID компании']
        }
        return headers_map.get(sheet_key, [])

    def add_queries(self, queries: List[Dict[str, Any]]) -> bool:
        """
        Добавить запросы в лист ЗАПРОСЫ

        Args:
            queries: Список запросов с полями: тема, город, запрос, статус

        Returns:
            True если успешно
        """
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.sheets['queries'])

            # Получаем следующий ID
            next_id = len(df) + 1

            # Подготавливаем данные
            tz = pytz.timezone(TIMEZONE)
            now = datetime.now(tz).strftime('%Y-%m-%d %H:%M:%S')

            new_rows = []
            for query in queries:
                row = {
                    'ID': next_id,
                    'Дата создания': now,
                    'Тема': query.get('тема', ''),
                    'Город': query.get('город', ''),
                    'Запрос': query.get('запрос', ''),
                    'Статус': query.get('статус', 0),
                    'Дата обработки': ''
                }
                new_rows.append(row)
                next_id += 1

            # Добавляем в DataFrame
            new_df = pd.DataFrame(new_rows)
            df = pd.concat([df, new_df], ignore_index=True)

            # Сохраняем
            with pd.ExcelWriter(self.excel_file, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
                df.to_excel(writer, sheet_name=self.sheets['queries'], index=False)

            logger.info(f"Добавлено {len(new_rows)} запросов в Excel")
            return True

        except Exception as e:
            logger.error(f"Ошибка при добавлении запросов: {e}")
            return False

    def get_queries_with_status(self, status: int = 0) -> List[Dict[str, Any]]:
        """
        Получить запросы с определённым статусом

        Args:
            status: Статус запросов (0 = не обработан, 1 = обработан)

        Returns:
            Список словарей с данными запросов
        """
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.sheets['queries'])

            # Фильтруем по статусу
            filtered = df[df['Статус'] == status]

            # Конвертируем в список словарей
            result = filtered.to_dict('records')

            logger.info(f"Найдено {len(result)} запросов со статусом {status}")
            return result

        except Exception as e:
            logger.error(f"Ошибка при получении запросов: {e}")
            return []

    def add_sites(self, sites: List[Dict[str, Any]]) -> bool:
        """Добавить сайты компаний в лист САЙТЫ"""
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.sheets['sites'])
            next_id = len(df) + 1

            tz = pytz.timezone(TIMEZONE)
            now = datetime.now(tz).strftime('%Y-%m-%d %H:%M:%S')

            new_rows = []
            for site in sites:
                row = {
                    'ID': next_id,
                    'Дата добавления': now,
                    'Название компании': site.get('название', ''),
                    'Сайт': site.get('сайт', ''),
                    'Телефон': site.get('телефон', ''),
                    'Адрес': site.get('адрес', ''),
                    'Город': site.get('город', ''),
                    'Категория': site.get('категория', ''),
                    'Статус сбора данных': site.get('статус', 0),
                    'ID запроса': site.get('id_запроса', '')
                }
                new_rows.append(row)
                next_id += 1

            new_df = pd.DataFrame(new_rows)
            df = pd.concat([df, new_df], ignore_index=True)

            with pd.ExcelWriter(self.excel_file, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
                df.to_excel(writer, sheet_name=self.sheets['sites'], index=False)

            logger.info(f"Добавлено {len(new_rows)} сайтов в Excel")
            return True

        except Exception as e:
            logger.error(f"Ошибка при добавлении сайтов: {e}")
            return False

    def get_sites_with_status(self, status: int = 0) -> List[Dict[str, Any]]:
        """Получить сайты с определённым статусом сбора данных"""
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.sheets['sites'])
            filtered = df[df['Статус сбора данных'] == status]
            result = filtered.to_dict('records')

            logger.info(f"Найдено {len(result)} сайтов со статусом {status}")
            return result

        except Exception as e:
            logger.error(f"Ошибка при получении сайтов: {e}")
            return []

    def add_companies(self, companies: List[Dict[str, Any]]) -> bool:
        """Добавить информацию о компаниях в лист КОМПАНИИ"""
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.sheets['companies'])
            next_id = len(df) + 1

            tz = pytz.timezone(TIMEZONE)
            now = datetime.now(tz).strftime('%Y-%m-%d %H:%M:%S')

            new_rows = []
            for company in companies:
                row = {
                    'ID': next_id,
                    'Дата сбора': now,
                    'Название компании': company.get('название', ''),
                    'Сайт': company.get('сайт', ''),
                    'Email': company.get('email', ''),
                    'Телефон': company.get('телефон', ''),
                    'Адрес': company.get('адрес', ''),
                    'Описание': company.get('описание', ''),
                    'Услуги': company.get('услуги', ''),
                    'Контактное лицо': company.get('контактное_лицо', ''),
                    'Соц. сети': company.get('соцсети', ''),
                    'ID сайта': company.get('id_сайта', '')
                }
                new_rows.append(row)
                next_id += 1

            new_df = pd.DataFrame(new_rows)
            df = pd.concat([df, new_df], ignore_index=True)

            with pd.ExcelWriter(self.excel_file, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
                df.to_excel(writer, sheet_name=self.sheets['companies'], index=False)

            logger.info(f"Добавлено {len(new_rows)} компаний в Excel")
            return True

        except Exception as e:
            logger.error(f"Ошибка при добавлении компаний: {e}")
            return False

    def add_letters(self, letters: List[Dict[str, Any]]) -> bool:
        """Добавить черновики писем в лист ПИСЬМА"""
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.sheets['letters'])
            next_id = len(df) + 1

            tz = pytz.timezone(TIMEZONE)
            now = datetime.now(tz).strftime('%Y-%m-%d %H:%M:%S')

            new_rows = []
            for letter in letters:
                row = {
                    'ID': next_id,
                    'Дата создания': now,
                    'Название компании': letter.get('название', ''),
                    'Email': letter.get('email', ''),
                    'Тема письма': letter.get('тема', ''),
                    'Текст письма': letter.get('текст', ''),
                    'Статус отправки': letter.get('статус', 0),
                    'ID компании': letter.get('id_компании', '')
                }
                new_rows.append(row)
                next_id += 1

            new_df = pd.DataFrame(new_rows)
            df = pd.concat([df, new_df], ignore_index=True)

            with pd.ExcelWriter(self.excel_file, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
                df.to_excel(writer, sheet_name=self.sheets['letters'], index=False)

            logger.info(f"Добавлено {len(new_rows)} писем в Excel")
            return True

        except Exception as e:
            logger.error(f"Ошибка при добавлении писем: {e}")
            return False

    def update_query_status(self, query_id: int, status: int = 1):
        """Обновить статус запроса"""
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.sheets['queries'])

            tz = pytz.timezone(TIMEZONE)
            now = datetime.now(tz).strftime('%Y-%m-%d %H:%M:%S')

            # Обновляем статус
            df.loc[df['ID'] == query_id, 'Статус'] = status
            df.loc[df['ID'] == query_id, 'Дата обработки'] = now

            with pd.ExcelWriter(self.excel_file, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
                df.to_excel(writer, sheet_name=self.sheets['queries'], index=False)

            logger.info(f"Обновлён статус запроса {query_id} на {status}")

        except Exception as e:
            logger.error(f"Ошибка при обновлении статуса запроса: {e}")

    def update_site_status(self, site_id: int, status: int = 1):
        """Обновить статус сбора данных по сайту"""
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.sheets['sites'])

            df.loc[df['ID'] == site_id, 'Статус сбора данных'] = status

            with pd.ExcelWriter(self.excel_file, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
                df.to_excel(writer, sheet_name=self.sheets['sites'], index=False)

            logger.info(f"Обновлён статус сайта {site_id} на {status}")

        except Exception as e:
            logger.error(f"Ошибка при обновлении статуса сайта: {e}")

    def get_all_companies(self, limit: int = None) -> List[Dict[str, Any]]:
        """Получить все компании из Excel"""
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.sheets['companies'])

            if limit:
                df = df.head(limit)

            return df.to_dict('records')

        except Exception as e:
            logger.error(f"Ошибка при получении компаний: {e}")
            return []
